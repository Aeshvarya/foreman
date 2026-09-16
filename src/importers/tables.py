"""Reading the tables people actually export.

A P6 activity export, a P&D tracker and a procurement log all arrive as CSV or
Excel, and none of them agree on column names or on where the header row is:
trackers carry a title row or three, and P6's own Excel export has one row of
field names followed by a second row of display names. So each column is found
by name -- from a list of aliases per field -- in whichever of the first rows
matches best, and values are parsed the forgiving way: P6 marks actual dates
with " A" and constrained ones with "*", and Excel hands over real datetimes or
serial numbers instead of text.
"""
from __future__ import annotations

import csv
import io
import re
import sys
from datetime import date, datetime, timedelta
from typing import Iterable

csv.field_size_limit(sys.maxsize)   # exports can carry very wide columns (embeddings)

HEADER_SCAN_ROWS = 15


class ImportProblem(ValueError):
    """Something about an uploaded file that a person has to fix. The message
    says what, in words they can act on."""


def decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def norm_key(s) -> str:
    """'Original Duration(d)' -> 'original duration d', for matching headers."""
    return re.sub(r"[^a-z0-9#]+", " ", str(s if s is not None else "").lower()).strip()


def _sheets(filename: str, data: bytes) -> Iterable[Iterable[list]]:
    name = (filename or "").lower()
    if name.endswith(".xls"):
        raise ImportProblem(f"{filename}: old .xls files aren't supported — save it as .xlsx or CSV.")
    if name.endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
        except ImportError as e:
            raise ImportProblem("Excel files need openpyxl on the server — or save the sheet as CSV.") from e
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            sheets = [[list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets]
        finally:
            wb.close()
        yield from sheets
        return
    text = decode(data)
    first = text[:4096].split("\n", 1)[0]
    delim = max(("\t", ",", ";", "|"), key=first.count)
    yield csv.reader(io.StringIO(text), delimiter=delim if first.count(delim) else ",")


def _match(keys: list[str], wanted: dict[str, list[str]]) -> dict[str, int]:
    found: dict[str, int] = {}
    used: set[int] = set()
    for field, aliases in wanted.items():
        for alias in aliases:
            j = next((i for i, k in enumerate(keys) if k == alias and i not in used), None)
            if j is not None:
                found[field] = j
                used.add(j)
                break
    return found


def read_table(filename: str, data: bytes, fields: dict[str, list[str]], *,
               required: list[str], min_hits: int = 2,
               overrides: dict[str, str] | None = None) -> tuple[list[dict], dict[str, str]]:
    """Rows as {field: raw cell} dicts, plus {field: the header it came from}.

    `fields` maps each field to its accepted column names, most specific first.
    `overrides` pins a field to one exact column name.
    """
    overrides = {k: v for k, v in (overrides or {}).items() if v}
    wanted = {f: ([norm_key(overrides[f])] if f in overrides else []) + [norm_key(a) for a in aliases]
              for f, aliases in fields.items()}
    problem = None
    for sheet in _sheets(filename, data):
        it = iter(sheet)
        head: list[list] = []
        for row in it:
            head.append(list(row))
            if len(head) >= HEADER_SCAN_ROWS:
                break
        scores = [_match([norm_key(c) for c in row], wanted) for row in head]
        if not scores:
            continue
        best_i = max(range(len(scores)), key=lambda i: (len(scores[i]), -i))
        colmap = scores[best_i]
        missing = [f for f in required if f not in colmap]
        if len(colmap) < min_hits or missing:
            want = ", ".join(f"'{fields[f][0]}'" for f in (missing or required))
            problem = (f"{filename}: couldn't find the header row. It needs columns like {want} "
                       f"(found: {', '.join(str(c) for c in head[best_i][:12] if c) or 'nothing'}).")
            continue
        labels = {f: str(head[best_i][j]) for f, j in colmap.items()}
        rows: list[dict] = []

        def keep(row: list, near_header: bool) -> None:
            if not any(c is not None and str(c).strip() for c in row):
                return
            # P6's Excel export repeats the header as display names on the next row
            if near_header and len(_match([norm_key(c) for c in row], wanted)) >= max(min_hits, len(colmap) - 1):
                return
            rows.append({f: (row[j] if j < len(row) else None) for f, j in colmap.items()})

        for row in head[best_i + 1:]:
            keep(row, True)
        for row in it:
            keep(list(row), False)
        return rows, labels
    raise ImportProblem(problem or f"{filename}: that file looks empty.")


# ------------------------------------------------------------ cell parsing
_DATE_FORMATS = ("%Y-%m-%d", "%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%m/%d/%Y", "%m/%d/%Y %H:%M",
                 "%m/%d/%y", "%d/%m/%Y", "%d-%b-%y", "%d-%b-%Y", "%d-%b-%y %H:%M", "%d-%b-%Y %H:%M",
                 "%b %d, %Y", "%d %b %Y", "%d %B %Y", "%B %d, %Y", "%d.%m.%Y")


def to_datetime(v) -> datetime | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v
    if isinstance(v, date):
        return datetime(v.year, v.month, v.day)
    if isinstance(v, (int, float)):                 # an Excel serial day number
        return datetime(1899, 12, 30) + timedelta(days=float(v)) if 20000 < v < 80000 else None
    s = str(v).strip()
    s = re.sub(r"\s*[A*]+$", "", s).strip()          # P6: " A" = actual, "*" = constrained
    if re.search(r"\d:\d\d", s):                       # drop fractional seconds / time zones
        s = re.sub(r"(\.\d+)?(Z|[+-]\d\d(:?\d\d)?)?$", "", s.replace("T", " "))
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def to_date(v) -> date | None:
    dt = to_datetime(v)
    return dt.date() if dt else None


def to_num(v) -> float | None:
    """'12', '12.5', '12d', '1,200', 12 -> a number; blanks and words -> None."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    m = re.match(r"^\s*(-?\d+(?:\.\d+)?)", str(v).replace(",", ""))
    return float(m.group(1)) if m else None


def to_bool(v) -> bool:
    return str(v).strip().lower() in ("true", "t", "y", "yes", "1", "x")


def text(v) -> str:
    return "" if v is None else str(v).strip()
